"""
Write the documents of the FICTIONAL example artists into examples/artists/.

These personas are invented for development and tests. They are not real people,
so the knowledge-base code can be built and tested without anyone's real CV.
Each one covers different file formats, so every reader is exercised:

    ilka_varga/      cv.pdf (plain-text headings)    statement.docx (Word heading styles)
    tomas_ferreira/  cv.docx (exhibitions as a table) statement.md    portfolio.txt

    uv run python scripts/make_example_artists.py

The PDF is written by a tiny writer below (one font, text only), so no PDF
library is needed just for test files.
"""

import os
import textwrap

from docx import Document

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(REPO_ROOT, "examples", "artists")

ILKA_CV = """\
ILKA VARGA
Painter and printmaker, based in Vienna

PERSONAL DETAILS
Born 12 March 1994 in Debrecen, Hungary
Citizenship: Hungarian
Lives and works in Vienna, Austria, since 2018
Email: ilka.varga@example.org

EDUCATION
2019-2021 MFA in Painting and Graphic Arts, Academy of Fine Arts Vienna, Austria
2013-2017 BA in Painting, Hungarian University of Fine Arts, Budapest, Hungary

SOLO EXHIBITIONS
2025 Low Tide Archive, Galerie Nord, Vienna
2023 Paper Weather, Studio Kapu, Budapest

GROUP EXHIBITIONS
2024 New Positions in Print, Künstlerhaus Graz, Austria
2022 Surface Tension, Kunstraum Lakeside, Klagenfurt, Austria
2021 Diploma Exhibition, Academy of Fine Arts Vienna

RESIDENCIES
2024 Two-month printmaking residency, Atelier Lofoten, Norway
2022 Studio residency, Cité internationale des arts, Paris, France

AWARDS AND GRANTS
2025 Austrian federal working grant for visual arts
2021 Academy prize for printmaking, Academy of Fine Arts Vienna

LANGUAGES
Hungarian (native), German (fluent), English (fluent)
"""

ILKA_STATEMENT = [
    ("Title", "Artist statement"),
    ("Normal", "Ilka Varga"),
    ("Heading 1", "Practice"),
    ("Normal",
     "I work between painting and printmaking, building large layered works on paper from monotypes, "
     "etchings and thin washes of oil paint. My subject is weather as it is recorded rather than felt: "
     "tide tables, flood marks, the faded logbooks of river stations along the Danube. I copy, enlarge "
     "and overprint these records until the numbers become landscape. The work is slow and physical; "
     "each sheet passes through the press many times, and the errors of registration are kept, not "
     "corrected."),
    ("Heading 1", "Current project"),
    ("Normal",
     "Low Tide Archive (2024-) is a series of twenty prints made from the water-level registers of five "
     "Danube stations between 1950 and today. I am looking for residencies with access to an etching "
     "press and to a local archive, to extend the project to a coastline outside Central Europe."),
    ("Heading 1", "Methods"),
    ("Normal", "Monotype, etching, aquatint, oil on paper, archival research, large-format printing."),
]

TOMAS_CV_INTRO = [
    ("Title", "Tomás Ferreira"),
    ("Normal", "Sound artist — installations, field recording, live performance"),
    ("Heading 1", "Personal details"),
    ("Normal", "Born 4 July 1988 in Porto Alegre, Brazil"),
    ("Normal", "Nationality: Portuguese and Brazilian (dual citizenship)"),
    ("Normal", "Based in Lisbon, Portugal"),
    ("Normal", "Works as an individual artist and as a founding member of the collective Ruído Comum"),
    ("Heading 1", "Education"),
    ("Normal", "2006–2010 BA in Music Technology, Universidade Federal do Rio Grande do Sul, Brazil"),
    ("Normal", "Not currently enrolled in any study programme"),
    ("Heading 1", "Exhibitions and performances"),
]

TOMAS_CV_TABLE = [
    ("Year", "Work", "Venue"),
    ("2025", "Tidal Bodies (8-channel installation)", "Galeria Zé dos Bois, Lisbon"),
    ("2024", "Listening to the Tagus (live set)", "Festival Rescaldo, Lisbon"),
    ("2023", "Underwater Choir (installation)", "Kunsthall Bergen, Norway"),
    ("2021", "Rain Score (performance)", "Casa da Música, Porto"),
]

TOMAS_CV_OUTRO = [
    ("Heading 1", "Residencies"),
    ("Normal", "2023 Sound residency, Lydgalleriet, Bergen, Norway"),
    ("Normal", "2019 Field-recording residency, Ilha do Pico, Azores"),
    ("Heading 1", "Grants"),
    ("Normal", "2024 Direção-Geral das Artes project grant for Tidal Bodies"),
]

TOMAS_STATEMENT = """\
# Statement

I make sound installations and live performances from field recordings of water: rivers, harbours,
rain on different roofs, the inside of pipes. I record with hydrophones and contact microphones and
play the material back through many small speakers placed in the space, so that the listener walks
through the sound rather than in front of it.

# Themes

My recent work is about rivers that cities have hidden or buried. I am interested in residencies by
water, in commissions for public space, and in collaborations with ecologists and hydrologists.

# Collaborations

With the collective Ruído Comum I organise free listening walks in Lisbon neighbourhoods.
"""

TOMAS_PORTFOLIO = """\
SELECTED WORKS

Tidal Bodies, 2025
8-channel sound installation, 40 minutes, loop. Hydrophone recordings from the Tagus estuary at
high and low tide, played through speakers hung at different heights.

Underwater Choir, 2023
Installation for 24 small speakers and water tanks. Made during a residency in Bergen with
recordings from the city harbour.

Rain Score, 2021
Performance for amplified roofs. Four performers play rainwater falling on zinc, wood and slate.
"""


# -- a tiny text-only PDF writer --------------------------------------------------------------------------

def _pdf_string(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    return b"(" + escaped.encode("cp1252", errors="replace") + b")"


def write_pdf(path: str, text: str, lines_per_page: int = 52) -> None:
    lines = []
    for line in text.splitlines():
        lines += textwrap.wrap(line, 95) or [""]
    pages = [lines[i:i + lines_per_page] for i in range(0, len(lines), lines_per_page)]

    # object numbers: 1 catalog, 2 page tree, 3 font, then a (page, content) pair per page
    objects = {1: b"<< /Type /Catalog /Pages 2 0 R >>",
               3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"}
    kids = []
    for n, page_lines in enumerate(pages):
        page_obj, content_obj = 4 + 2 * n, 5 + 2 * n
        kids.append(f"{page_obj} 0 R".encode())
        stream = b"BT /F1 10 Tf 14 TL 50 800 Td " + b" ".join(_pdf_string(l) + b" Tj T*" for l in page_lines) + b" ET"
        objects[content_obj] = b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
        objects[page_obj] = (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
                             b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % content_obj)
    objects[2] = b"<< /Type /Pages /Kids [" + b" ".join(kids) + b"] /Count %d >>" % len(pages)

    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for number in sorted(objects):
        offsets[number] = len(out)
        out += b"%d 0 obj\n" % number + objects[number] + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offsets[n] for n in sorted(objects))
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    with open(path, "wb") as f:
        f.write(out)


def write_docx(path: str, paragraphs: list[tuple[str, str]], table=None, after=()) -> None:
    document = Document()
    for style, text in paragraphs:
        document.add_paragraph(text, style=style)
    if table:
        grid = document.add_table(rows=0, cols=len(table[0]))
        for row in table:
            for cell, value in zip(grid.add_row().cells, row):
                cell.text = value
    for style, text in after:
        document.add_paragraph(text, style=style)
    document.save(path)


def main() -> None:
    ilka, tomas = os.path.join(OUT_DIR, "ilka_varga"), os.path.join(OUT_DIR, "tomas_ferreira")
    os.makedirs(ilka, exist_ok=True)
    os.makedirs(tomas, exist_ok=True)

    write_pdf(os.path.join(ilka, "cv.pdf"), ILKA_CV)
    write_docx(os.path.join(ilka, "statement.docx"), ILKA_STATEMENT)
    write_docx(os.path.join(tomas, "cv.docx"), TOMAS_CV_INTRO, TOMAS_CV_TABLE, TOMAS_CV_OUTRO)
    with open(os.path.join(tomas, "statement.md"), "w", encoding="utf-8") as f:
        f.write(TOMAS_STATEMENT)
    with open(os.path.join(tomas, "portfolio.txt"), "w", encoding="utf-8") as f:
        f.write(TOMAS_PORTFOLIO)
    print(f"wrote example artists to {os.path.relpath(OUT_DIR, REPO_ROOT)}/")


if __name__ == "__main__":
    main()
