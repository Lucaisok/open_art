"""
OpenArt API — an artist's application to one call (workflow.MD, "Step 6 — Application agent").

    GET  /api/applications/{opportunity_id}                         the checklist with the artist's
                                                                    drafts, ticks and documents
    PUT  /api/applications/{opportunity_id}/items/{key}             save a draft (back to "draft"), a tick,
                                                                    or "Not needed" (skipped)
    POST /api/applications/{opportunity_id}/items/{key}/confirm     "I have reviewed this"
    PUT  /api/applications/{opportunity_id}/projects                the projects chosen for the application
    DELETE /api/applications/{opportunity_id}/projects              ...back to the suggested ones
    PUT  /api/applications/{opportunity_id}/items/{key}/projects    ...or for one text only (null = the application's)
    PUT  /api/applications/{opportunity_id}/items/{key}/images      the images chosen for an image item
    POST /api/applications/{opportunity_id}/texts                   add a text of the artist's own, with a title
    PUT  /api/applications/{opportunity_id}/texts/{key}             ...rename it
    DELETE /api/applications/{opportunity_id}/texts/{key}           ...remove it, with its chat
    GET  /api/applications/{opportunity_id}/items/{key}/messages    the chat about one item
    POST /api/applications/{opportunity_id}/items/{key}/messages    a message to the agent: its reply
                                                                    and a new draft (src/agent/drafting.py)
    GET  /api/applications/{opportunity_id}/download                the complete application as a zip: the
                                                                    reviewed texts + the requested documents

What the call asks for comes from the corpus every time (src/agent/checklist.py); only the artist's
own work is stored, including the texts they add themselves on any call (title set, key "t-...") (api/models.py: Application, ApplicationItem, AgentMessage), created on their first save.
Human in the loop: only the confirm route sets "confirmed", and only on a draft the artist can see;
any later change to the text, the agent's included, sets it back to "draft". The agent's model call
(OpenAI) is the only thing that leaves the server, and nothing is ever sent to the call itself.
"""

import io
import os
import re
import secrets
import uuid
import zipfile
from datetime import date, datetime, timezone
from functools import lru_cache
from typing import Literal

from docx import Document as DocxDocument
from fastapi import APIRouter, HTTPException
import numpy as np
from fastapi.responses import Response
from openai import OpenAI, OpenAIError
from pydantic import BaseModel, Field
from sqlalchemy import delete, select

from api.auth import DB, CurrentUser
from api import discover   # discover.get_matcher(), not imported by name: the tests replace it
from api.documents import stored_path
from api.knowledge_base import PgKnowledgeBase
from api.images import ImageOut, image_out, image_path
from api.models import AgentMessage, Application, ApplicationItem, Document, Profile, Project, ProjectImage, SavedCall
from api.projects import can_propose, project_text
from src.agent.checklist import ChecklistItem, build_checklist
from src.agent.drafting import Passage, TurnContext, agent_turn
from src.eligibility.engine import Status
from src.matching.index import embed_query
from src.models.processed_opportunity import ProcessedOpportunity

router = APIRouter(prefix="/api/applications", tags=["applications"])

MAX_DRAFT_CHARS = 20_000   # far above any call's limit; only stops abuse


class Selection(BaseModel):
    """Which saved projects a text is about: the one being proposed, and past or related work."""
    proposing: uuid.UUID | None = None
    mention: list[uuid.UUID] = Field(default=[], max_length=50)


class ProjectChoiceOut(BaseModel):
    id: uuid.UUID
    title: str
    status: str | None          # past / current / planned
    category: str | None        # work / exhibition / performance / commission (None: saved before categories)
    years: str | None
    medium: str | None = None
    can_propose: bool = False   # a planned or current work (projects.can_propose): the one an application proposes
    suggested: bool             # in the suggested selection (closest to the call)
    images: list[ImageOut] = []  # its images, shown as thumbnails in the picker


class ItemOut(BaseModel):
    key: str
    kind: Literal["write", "gather"]
    source: Literal["question", "material", "artist"]   # "artist": a text the artist added
    title: str
    limit: str | None
    document: str | None        # "cv" / "statement" / "portfolio" that covers it (uploaded or not)
    # write: "empty" / "draft" / "confirmed"; gather: "uploaded" (its document is) / "chosen" (images
    # chosen for it) / "ticked" / "missing"
    state: Literal["empty", "draft", "confirmed", "uploaded", "chosen", "ticked", "missing"]
    images_wanted: bool = False         # asks for pictures of the work (checklist.asks_for_images)
    images: list[ImageOut] = []         # the ones chosen, in the artist's order
    draft: str | None
    projects: Selection | None = None   # this text's own choice; None = the application's
    skipped: bool = False               # the artist's "Not needed": out of the download and its rule


class ProjectImagesOut(BaseModel):
    """One saved project's images, to choose from for an image item."""
    project_id: uuid.UUID
    title: str
    years: str | None = None
    images: list[ImageOut]


class ImagesIn(BaseModel):
    ids: list[uuid.UUID] = Field(default=[], max_length=50)   # in the order chosen; [] = none


class ApplicationDocumentOut(BaseModel):
    kind: str
    file_name: str


class ApplicationOut(BaseModel):
    opportunity_id: str
    title: str
    organisation: str | None = None
    city: str | None = None
    countries: list[str] = []
    saved: bool = False               # in the artist's saved calls (api/saved.py)
    deadline: date | None
    source_url: str
    how_to_apply: str | None
    application_language: str | None
    submission_method: str | None
    verdict_status: Status            # the engine's, as on the opportunity page: never changed here
    to_draft: list[ItemOut]
    to_gather: list[ItemOut]
    documents: list[ApplicationDocumentOut]
    # what stands before the download ("2 texts to review"); [] = complete. Same rule as the download route
    missing: list[str] = []
    projects: list[ProjectChoiceOut] = []   # the artist's saved projects, closest to the call first
    selection: Selection = Selection()      # the application's choice (the suggested one until chosen)
    selection_chosen: bool = False          # False = still the suggestion
    # the saved projects' images, when an item asks for images (else [])
    image_choices: list[ProjectImagesOut] = []


class TextIn(BaseModel):
    """A text of the artist's own: its title ("Cover letter")."""
    title: str = Field(min_length=1, max_length=200)


class ItemIn(BaseModel):
    """Any of the fields. A draft only for an item to write, a tick only for one to gather; "Not
    needed" (skipped) for any item the call asks for (the artist's own texts are deleted instead)."""
    draft: str | None = Field(default=None, max_length=MAX_DRAFT_CHARS)
    ticked: bool | None = None
    skipped: bool | None = None


def _call(opportunity_id: str) -> ProcessedOpportunity:
    call = discover.get_matcher().get(opportunity_id)
    if call is None:
        raise HTTPException(404, "No such call.")
    return call


def _application(db, user_id: uuid.UUID, opportunity_id: str, create: bool = False) -> Application | None:
    application = db.scalar(select(Application).where(Application.user_id == user_id,
                                                      Application.opportunity_id == opportunity_id))
    if application is None and create:
        application = Application(user_id=user_id, opportunity_id=opportunity_id)
        db.add(application)
        db.flush()   # gives it its id
        # starting an application saves the call: the Studio lists it (api/saved.py)
        from api.saved import save_call   # api/saved.py imports this module
        save_call(db, user_id, discover.get_matcher().get(opportunity_id))
    return application


def own_texts(application: Application | None) -> list[ChecklistItem]:
    """The texts the artist added themselves, in the order they added them: items to write like the
    call's own, with the artist's title."""
    if application is None:
        return []
    rows = sorted((i for i in application.items if i.title), key=lambda i: i.created_at)
    return [ChecklistItem(key=i.key, kind="write", source="artist", title=i.title, limit=None, document=None)
            for i in rows]


def _all_items(call: ProcessedOpportunity, application: Application | None) -> list[ChecklistItem]:
    """What the call asks for, then the artist's own texts."""
    return build_checklist(call) + own_texts(application)


def _chosen(saved: ApplicationItem | None, images: dict[str, ProjectImage]) -> list[ProjectImage]:
    """The images chosen for an item that still exist and belong to a saved project."""
    return [images[i] for i in (saved.images or []) if i in images] if saved else []   # deleted since: left out


def _state(item: ChecklistItem, saved: ApplicationItem | None, uploaded: set[str], chosen: list) -> str:
    if item.kind == "write":
        return saved.status if saved and saved.status else "empty"
    if item.document in uploaded:
        return "uploaded"
    if chosen:
        return "chosen"
    return "ticked" if saved and saved.ticked else "missing"


def _item_out(item: ChecklistItem, saved: ApplicationItem | None, uploaded: set[str],
              images: dict[str, ProjectImage]) -> ItemOut:
    chosen = _chosen(saved, images) if item.images else []
    return ItemOut(key=item.key, kind=item.kind, source=item.source, title=item.title, limit=item.limit,
                   document=item.document, state=_state(item, saved, uploaded, chosen),
                   draft=saved.draft if saved else None,
                   projects=Selection(**saved.projects) if saved and saved.projects else None,
                   skipped=bool(saved and saved.skipped),
                   images_wanted=item.images, images=[image_out(i) for i in chosen])


def _project_images(db, user_id: uuid.UUID) -> dict[str, ProjectImage]:
    """{image id: image} of every image a saved project holds: what an image item can use."""
    rows = db.scalars(select(ProjectImage).where(ProjectImage.user_id == user_id, ProjectImage.project_id.is_not(None))
                      .order_by(ProjectImage.position, ProjectImage.created_at))
    return {str(i.id): i for i in rows}


def application_out(db, user_id: uuid.UUID, call: ProcessedOpportunity) -> ApplicationOut:
    documents = list(db.scalars(select(Document).where(Document.user_id == user_id).order_by(Document.kind)))
    uploaded = {d.kind for d in documents}
    application = _application(db, user_id, call.id)
    saved = {i.key: i for i in application.items} if application else {}
    images = _project_images(db, user_id)
    checklist = _all_items(call, application)
    items = [_item_out(i, saved.get(i.key), uploaded, images) for i in checklist]
    profile, _ = discover.saved_profile(db, user_id)
    city, countries = discover.place(call)
    out = ApplicationOut(
        opportunity_id=call.id, title=call.title_en, organisation=call.organisation, city=city, countries=countries,
        saved=db.get(SavedCall, (user_id, call.id)) is not None,
        deadline=call.deadline_date, source_url=call.source_url,
        how_to_apply=call.how_to_apply_en or call.how_to_apply,
        application_language=call.application_language_en or call.application_language,
        submission_method=call.submission_method,
        verdict_status=discover.get_matcher().engine.evaluate(profile, call).status,
        to_draft=[i for i in items if i.kind == "write"], to_gather=[i for i in items if i.kind == "gather"],
        documents=[ApplicationDocumentOut(kind=d.kind, file_name=d.file_name) for d in documents])
    out.missing = missing(out)
    projects = _saved_projects(db, user_id)
    suggested = suggested_projects(call, projects)
    by_id = {p.id: p for p in projects}
    ranked = [by_id[i] for i in ([suggested.proposing] if suggested.proposing else []) + suggested.mention]
    ranked += [p for p in projects if p not in ranked]
    picked = {suggested.proposing, *suggested.mention}
    by_project: dict[uuid.UUID, list[ImageOut]] = {}
    for image in images.values():
        by_project.setdefault(image.project_id, []).append(image_out(image))
    out.projects = [ProjectChoiceOut(id=p.id, title=p.title or "Untitled project", status=p.status,
                                     category=p.category, years=p.years, medium=p.medium, can_propose=can_propose(p),
                                     suggested=p.id in picked, images=by_project.get(p.id, [])) for p in ranked]
    out.selection_chosen = bool(application and application.projects is not None)
    out.selection = _clean(Selection(**application.projects), projects) if out.selection_chosen else suggested
    if any(i.images for i in checklist):
        out.image_choices = [ProjectImagesOut(project_id=p.id, title=p.title or "Untitled project", years=p.years,
                                              images=by_project[p.id]) for p in ranked if p.id in by_project]
    return out


def _item(db, user_id: uuid.UUID, call: ProcessedOpportunity, key: str) -> ChecklistItem:
    """One item of this artist's application: one the call asks for, or a text they added."""
    item = next((i for i in _all_items(call, _application(db, user_id, call.id)) if i.key == key), None)
    if item is None:
        raise HTTPException(404, "This application has no such item.")
    return item


def _saved_item(db, application: Application, key: str) -> ApplicationItem:
    saved = db.get(ApplicationItem, (application.id, key))
    if saved is None:
        saved = ApplicationItem(application_id=application.id, key=key, ticked=False)
        db.add(saved)
    return saved


@router.get("/{opportunity_id}")
def get_application(opportunity_id: str, user: CurrentUser, db: DB) -> ApplicationOut:
    return application_out(db, user.id, _call(opportunity_id))


@router.put("/{opportunity_id}/items/{key}")
def save_item(opportunity_id: str, key: str, body: ItemIn, user: CurrentUser, db: DB) -> ApplicationOut:
    call = _call(opportunity_id)
    item = _item(db, user.id, call, key)
    if body.draft is not None and item.kind != "write":
        raise HTTPException(422, "Only an item to write has a draft.")
    if body.ticked is not None and item.kind != "gather":
        raise HTTPException(422, "Only an item to gather can be ticked.")
    if body.skipped is not None and item.source == "artist":
        raise HTTPException(422, "Delete a text of your own instead.")
    saved = _saved_item(db, _application(db, user.id, call.id, create=True), key)
    if body.draft is not None and body.draft != saved.draft:
        saved.draft = body.draft
        saved.status = "draft" if body.draft.strip() else None   # a changed text needs a new review
    if body.ticked is not None:
        saved.ticked = body.ticked
    if body.skipped is not None:
        saved.skipped = body.skipped
    db.commit()
    return application_out(db, user.id, call)


@router.put("/{opportunity_id}/items/{key}/images")
def choose_images(opportunity_id: str, key: str, body: ImagesIn, user: CurrentUser, db: DB) -> ApplicationOut:
    """The images to send for an item asking for images: the artist's own choice from their saved
    projects' images, in their order. [] clears it."""
    call = _call(opportunity_id)
    item = _item(db, user.id, call, key)
    if not item.images:
        raise HTTPException(422, "This item doesn't ask for images.")
    images = _project_images(db, user.id)
    ids = list(dict.fromkeys(str(i) for i in body.ids))
    if any(i not in images for i in ids):
        raise HTTPException(422, "That isn't one of your projects' images.")
    saved = _saved_item(db, _application(db, user.id, call.id, create=True), key)
    saved.images = ids or None
    db.commit()
    return application_out(db, user.id, call)


@router.post("/{opportunity_id}/items/{key}/confirm")
def confirm_item(opportunity_id: str, key: str, user: CurrentUser, db: DB) -> ApplicationOut:
    call = _call(opportunity_id)
    _item(db, user.id, call, key)
    application = _application(db, user.id, call.id)
    saved = db.get(ApplicationItem, (application.id, key)) if application else None
    if saved is None or saved.status is None:
        raise HTTPException(409, "There is no draft to confirm yet.")
    saved.status = "confirmed"
    db.commit()
    return application_out(db, user.id, call)


# -- the artist's own texts (sub-step 8) ---------------------------------------------------------------

def _own_text(db, user_id: uuid.UUID, call: ProcessedOpportunity, key: str) -> ApplicationItem:
    """One of the texts this artist added on this call (never a checklist item: those can't be renamed)."""
    application = _application(db, user_id, call.id)
    saved = db.get(ApplicationItem, (application.id, key)) if application else None
    if saved is None or not saved.title:
        raise HTTPException(404, "There is no such text of yours.")
    return saved


@router.post("/{opportunity_id}/texts")
def add_text(opportunity_id: str, body: TextIn, user: CurrentUser, db: DB) -> ApplicationOut:
    """A text the artist wants to send, whether or not the call asks for it ("Cover letter"). Like the
    call's texts it is drafted, reviewed by the artist, and goes into the download."""
    call = _call(opportunity_id)
    title = body.title.strip()
    if not title:
        raise HTTPException(422, "Give the text a title.")
    application = _application(db, user.id, call.id, create=True)
    # "t-" + random: never the same as the call's "q-"/"m-" keys
    db.add(ApplicationItem(application_id=application.id, key=f"t-{secrets.token_hex(4)}", title=title, ticked=False))
    db.commit()
    db.refresh(application)   # its items now include the new text
    return application_out(db, user.id, call)


@router.put("/{opportunity_id}/texts/{key}")
def rename_text(opportunity_id: str, key: str, body: TextIn, user: CurrentUser, db: DB) -> ApplicationOut:
    call = _call(opportunity_id)
    title = body.title.strip()
    if not title:
        raise HTTPException(422, "Give the text a title.")
    _own_text(db, user.id, call, key).title = title
    db.commit()
    return application_out(db, user.id, call)


@router.delete("/{opportunity_id}/texts/{key}")
def delete_text(opportunity_id: str, key: str, user: CurrentUser, db: DB) -> ApplicationOut:
    """Removes the text, its draft and its chat with the assistant."""
    call = _call(opportunity_id)
    saved = _own_text(db, user.id, call, key)
    db.execute(delete(AgentMessage).where(AgentMessage.application_id == saved.application_id,
                                          AgentMessage.item_key == key))
    db.delete(saved)
    db.commit()
    db.expire_all()   # the application's items are read again without it
    return application_out(db, user.id, call)


# -- projects for the application (sub-step 7) ---------------------------------------------------------

SUGGESTED_MENTIONS = 3   # past or related projects pre-ticked
# project text -> its embedding: a project is embedded once per text, not on every page load
_project_vectors: dict[str, np.ndarray] = {}
MAX_CACHED_PROJECTS = 2000


def _saved_projects(db, user_id: uuid.UUID) -> list[Project]:
    """The artist's saved projects (not interests), oldest first: what can be chosen."""
    return list(db.scalars(select(Project).where(Project.user_id == user_id, Project.kind == "project")
                           .order_by(Project.created_at, Project.id)))


def suggested_projects(call: ProcessedOpportunity, projects: list[Project]) -> Selection:
    """A rule, no LLM: the projects closest to the call, embedded with Discover's model and compared
    with the call's own vector in the index. Proposing = the closest planned or current work (none
    if there is none); mention = the next closest three. The artist changes them freely."""
    matcher = discover.get_matcher()
    row = matcher._row.get(call.id)
    if row is None or not projects:
        return Selection()
    call_vector = matcher.index.vectors[row]
    if len(_project_vectors) > MAX_CACHED_PROJECTS:
        _project_vectors.clear()

    def closeness(p: Project) -> float:
        text = project_text(p)
        if text not in _project_vectors:
            _project_vectors[text] = embed_query(matcher.embedder, text)
        return float(_project_vectors[text] @ call_vector)

    ranked = sorted(projects, key=closeness, reverse=True)
    proposing = next((p for p in ranked if can_propose(p)), None)
    mention = [p.id for p in ranked if p is not proposing][:SUGGESTED_MENTIONS]
    return Selection(proposing=proposing.id if proposing else None, mention=mention)


def _clean(selection: Selection, projects: list[Project]) -> Selection:
    """Drops projects deleted since they were chosen; the proposed one is never also a mention."""
    ids = {p.id for p in projects}
    proposing = selection.proposing if selection.proposing in ids else None
    return Selection(proposing=proposing, mention=[i for i in dict.fromkeys(selection.mention)
                                                   if i in ids and i != proposing])


def _check_selection(selection: Selection, projects: list[Project]) -> None:
    ids = {p.id for p in projects}
    if (selection.proposing and selection.proposing not in ids) or any(i not in ids for i in selection.mention):
        raise HTTPException(422, "That isn't one of your projects.")


def effective_selection(db, user_id: uuid.UUID, call: ProcessedOpportunity, application: Application | None,
                        saved: ApplicationItem | None) -> Selection:
    """What a text uses: its own choice, else the application's, else the suggestion."""
    projects = _saved_projects(db, user_id)
    if saved and saved.projects is not None:
        return _clean(Selection(**saved.projects), projects)
    if application and application.projects is not None:
        return _clean(Selection(**application.projects), projects)
    return suggested_projects(call, projects)


class ItemProjectsIn(BaseModel):
    selection: Selection | None   # None: back to the application's choice


@router.put("/{opportunity_id}/projects")
def choose_projects(opportunity_id: str, body: Selection, user: CurrentUser, db: DB) -> ApplicationOut:
    call = _call(opportunity_id)
    _check_selection(body, _saved_projects(db, user.id))
    application = _application(db, user.id, call.id, create=True)
    application.projects = _clean(body, _saved_projects(db, user.id)).model_dump(mode="json")
    db.commit()
    return application_out(db, user.id, call)


@router.delete("/{opportunity_id}/projects")
def suggested_again(opportunity_id: str, user: CurrentUser, db: DB) -> ApplicationOut:
    """Back to the suggested projects (the closest to the call): the artist's own choice is forgotten."""
    call = _call(opportunity_id)
    application = _application(db, user.id, call.id)
    if application is not None:
        application.projects = None
        db.commit()
    return application_out(db, user.id, call)


@router.put("/{opportunity_id}/items/{key}/projects")
def choose_item_projects(opportunity_id: str, key: str, body: ItemProjectsIn, user: CurrentUser,
                         db: DB) -> ApplicationOut:
    call = _call(opportunity_id)
    if _item(db, user.id, call, key).kind != "write":
        raise HTTPException(422, "Only a text to write has projects.")
    if body.selection is not None:
        _check_selection(body.selection, _saved_projects(db, user.id))
    saved = _saved_item(db, _application(db, user.id, call.id, create=True), key)
    saved.projects = (_clean(body.selection, _saved_projects(db, user.id)).model_dump(mode="json")
                      if body.selection is not None else None)
    db.commit()
    return application_out(db, user.id, call)


# -- the agent (sub-step 4) ----------------------------------------------------------------------------

PASSAGES = 6             # the artist's passages closest to the item
STATEMENT_PASSAGES = 4   # plus, when adapting the uploaded statement, its own closest passages


@lru_cache
def get_openai() -> OpenAI:
    """One client per process; reads OPENAI_API_KEY (docker-compose.yml). The tests replace it."""
    return OpenAI()


class MessageOut(BaseModel):
    role: Literal["artist", "agent"]
    content: str
    sources: list[str]          # the agent's: citations of the passages it relied on
    wrote_words: int | None = None   # the agent's: it rewrote the draft, now this many words
    created_at: datetime


class MessageIn(BaseModel):
    message: str = Field(min_length=1, max_length=4000)
    # the text in the draft box now, saved or not: the agent revises what the artist sees
    draft: str | None = Field(default=None, max_length=MAX_DRAFT_CHARS)


class TurnOut(BaseModel):
    application: ApplicationOut
    messages: list[MessageOut]


def _messages(db, application: Application | None, key: str) -> list[AgentMessage]:
    if application is None:
        return []
    return list(db.scalars(select(AgentMessage).where(AgentMessage.application_id == application.id,
                                                      AgentMessage.item_key == key)
                           .order_by(AgentMessage.created_at, AgentMessage.id)))


def _message_out(m: AgentMessage) -> MessageOut:
    return MessageOut(role=m.role, content=m.content, sources=m.sources, wrote_words=m.wrote_words,
                      created_at=m.created_at)


def _passages(db, user_id: uuid.UUID, query: str, statement_file: str | None) -> list[Passage]:
    """The artist's passages closest to the item, plus the statement's own when adapting it."""
    kb = PgKnowledgeBase(db, user_id)
    try:
        found = kb.retrieve(query, k=PASSAGES)
        if statement_file:
            found += kb.retrieve(query, k=STATEMENT_PASSAGES, documents=[statement_file])
    except RuntimeError:
        raise HTTPException(409, "Your documents need to be uploaded again before the assistant can use them.")
    unique = {p.chunk.chunk_id: p for p in found}   # a statement passage can be in both lists
    return [Passage(id=f"P{i}", citation=p.citation, text=p.chunk.text)
            for i, p in enumerate(unique.values(), start=1)]


def _profile(db, user_id: uuid.UUID) -> dict[str, object]:
    row = db.get(Profile, user_id)
    return {k: v for k, v in (row.values if row else {}).items() if v not in (None, [], "")}


def _project_line(p: Project) -> str:
    when = ", ".join(x for x in (p.status, p.years, p.medium) if x)
    return f"{p.title or 'Untitled project'} ({when}): {p.summary or ''}".strip()


def _projects(db, user_id: uuid.UUID, selection: Selection) -> list[str]:
    """The chosen projects only (sub-step 7), the proposed one marked, and every interest."""
    by_id = {p.id: p for p in _saved_projects(db, user_id)}
    if selection.proposing:
        lines = [f"PROPOSING (the project this application is about): {_project_line(by_id[selection.proposing])}"]
    else:
        lines = ["No project is marked as the one being proposed."]
    lines += [f"Past or related work: {_project_line(by_id[i])}" for i in selection.mention]
    lines += [f"Interest: {p.summary}" for p in db.scalars(select(Project).where(Project.user_id == user_id,
                                                                              Project.kind == "interest"))]
    return lines


@router.get("/{opportunity_id}/items/{key}/messages")
def get_messages(opportunity_id: str, key: str, user: CurrentUser, db: DB) -> list[MessageOut]:
    call = _call(opportunity_id)
    _item(db, user.id, call, key)
    return [_message_out(m) for m in _messages(db, _application(db, user.id, call.id), key)]


@router.post("/{opportunity_id}/items/{key}/messages")
def send_message(opportunity_id: str, key: str, body: MessageIn, user: CurrentUser, db: DB) -> TurnOut:
    call = _call(opportunity_id)
    item = _item(db, user.id, call, key)
    if item.kind != "write":
        raise HTTPException(422, "The assistant only drafts items to write.")
    application = _application(db, user.id, call.id, create=True)
    saved = _saved_item(db, application, key)
    history = _messages(db, application, key)

    # the application's memory: the other items' current drafts, by their titles
    titles = {i.key: i.title for i in _all_items(call, application)}
    other_drafts = [(titles[i.key], i.draft) for i in application.items   # not the ones marked "Not needed"
                    if i.key != key and i.key in titles and i.draft and i.draft.strip() and not i.skipped]
    statement = db.scalar(select(Document).where(Document.user_id == user.id, Document.kind == "statement"))
    adapt = item.document == "statement" and statement is not None
    passages = _passages(db, user.id, f"{item.title}. {call.title_en}", statement.file_name if adapt else None)

    context = TurnContext(
        item_title=item.title, item_limit=item.limit,
        call={"Title": call.title_en, "Organisation": call.organisation, "Description": call.description_en,
              "Project requirements": call.project_requirements_en or call.project_requirements,
              "Selection criteria": call.selection_criteria_en or call.selection_criteria,
              "Application language": call.application_language_en or call.application_language},
        passages=passages, profile=_profile(db, user.id), projects=_projects(db, user.id, effective_selection(db, user.id, call, application, saved)),
        other_drafts=other_drafts, history=[(m.role, m.content) for m in history],
        current_draft=body.draft if body.draft is not None else saved.draft,
        message=body.message, adapt_statement=adapt)
    # times set here, not by the database: both rows are saved in one transaction, where now() is one instant
    asked_at = datetime.now(timezone.utc)
    try:
        answer = agent_turn(get_openai(), context)
    except OpenAIError:
        raise HTTPException(503, "The assistant isn't available right now. Please try again in a moment.")

    by_id = {p.id: p.citation for p in passages}
    sources = list(dict.fromkeys(by_id[i] for i in answer.used_passages if i in by_id))   # unique, in order
    db.add(AgentMessage(application_id=application.id, item_key=key, role="artist", content=body.message,
                        created_at=asked_at))
    previous = body.draft if body.draft is not None else saved.draft
    rewrote = bool(answer.draft.strip()) and answer.draft != (previous or "")
    db.add(AgentMessage(application_id=application.id, item_key=key, role="agent", content=answer.reply,
                        sources=sources, wrote_words=len(answer.draft.split()) if rewrote else None,
                        created_at=datetime.now(timezone.utc)))
    saved.draft = answer.draft
    saved.status = "draft" if answer.draft.strip() else None   # the artist reviews it: never "confirmed" here
    db.commit()
    return TurnOut(application=application_out(db, user.id, call),
                   messages=[_message_out(m) for m in _messages(db, application, key)])


# -- download (sub-step 6) -----------------------------------------------------------------------------

def without_skipped(application: ApplicationOut) -> ApplicationOut:
    """The application as it is sent: without the items the artist marked "Not needed"."""
    return application.model_copy(update={"to_draft": [i for i in application.to_draft if not i.skipped],
                                          "to_gather": [i for i in application.to_gather if not i.skipped]})


def missing(application: ApplicationOut) -> list[str]:
    """What still stands between the artist and a complete application: texts not reviewed, items
    neither uploaded, nor given images, nor ticked, or a zip with no file of the artist's in it
    (no text, no document, no image). Items marked "Not needed" don't count. [] = complete."""
    application = without_skipped(application)
    to_review = sum(i.state != "confirmed" for i in application.to_draft)
    to_gather = sum(i.state not in ("uploaded", "chosen", "ticked") for i in application.to_gather)
    out = []
    if to_review:
        out.append(f"{to_review} text{'s' if to_review > 1 else ''} to review")
    if to_gather:
        out.append(f"{to_gather} item{'s' if to_gather > 1 else ''} to gather")
    # a zip with only application.docx and the checklist holds nothing of the artist's: no download
    files = len(application.to_draft) + sum(i.state in ("uploaded", "chosen") for i in application.to_gather)
    if not out and files == 0:
        out.append("add a text (the zip would be empty)")
    return out


def _call_lines(call: ProcessedOpportunity, application: ApplicationOut) -> list[str]:
    lines = [call.title_en, f"Call: {call.source_url}"]
    if application.deadline:
        lines.append(f"Deadline: {application.deadline.isoformat()}")
    if application.submission_method:
        lines.append(f"Submit by: {application.submission_method}")
    return lines


def _image_folder(n: int, item: ItemOut) -> str:
    """Where an image item's images go in the zip: "images/01 Images of previous works"."""
    return f"images/{n:02d} {_safe_name(item.title)}"


def _to_send(application: ApplicationOut) -> list[str]:
    def where(n: int, i: ItemOut) -> str:
        if i.state == "uploaded":
            return f": in this folder ({i.document})"
        if i.state == "chosen":
            return f": {len(i.images)} image{'s' if len(i.images) > 1 else ''} in {_image_folder(n, i)}"
        return ": you have it"
    return [i.title + (f" ({i.limit})" if i.limit else "") + where(n, i)
            for n, i in enumerate(application.to_gather, start=1)]


def checklist_text(call: ProcessedOpportunity, application: ApplicationOut, text_files: list[str]) -> str:
    """checklist.txt: the call, and every file or thing that goes into the application."""
    lines = _call_lines(call, application) + ["", "Texts (reviewed by you):"]
    lines += [f"- {name}" for name in text_files]
    if application.to_gather:
        lines += ["", "To send with them:"] + [f"- {line}" for line in _to_send(application)]
    return "\n".join(lines) + "\n"


def application_docx(call: ProcessedOpportunity, application: ApplicationOut) -> bytes:
    """application.docx: every reviewed text under its own heading, for calls that want a document."""
    document = DocxDocument()
    document.add_heading(call.title_en, level=1)
    for line in _call_lines(call, application)[1:]:
        document.add_paragraph(line)
    for item in application.to_draft:
        document.add_heading(item.title + (f" ({item.limit})" if item.limit else ""), level=2)
        for paragraph in re.split(r"\n\s*\n", (item.draft or "").strip()):   # the text's own paragraphs
            document.add_paragraph(paragraph.strip())
    if application.to_gather:
        document.add_heading("To send with it", level=2)
        for line in _to_send(application):
            document.add_paragraph(line, style="List Bullet")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _safe_name(text: str) -> str:
    """A file name from free text: letters (accents too), digits, spaces and dashes only."""
    return re.sub(r"\s+", " ", re.sub(r"[^\w\- ]+", " ", text)).strip()[:60].strip() or "application"


@router.get("/{opportunity_id}/download")
def download(opportunity_id: str, user: CurrentUser, db: DB) -> Response:
    """The complete application as a zip. Refused until every text is reviewed and every item is
    uploaded or ticked: the review is enforced here, not only on the page. Holds each text as a .txt
    (to paste into forms), all of them in application.docx (to attach), checklist.txt, and for each item
    covered by an upload that original file (only the documents the call asks for)."""
    call = _call(opportunity_id)
    application = without_skipped(application_out(db, user.id, call))   # "Not needed" stays out of the zip
    gaps = missing(application)
    if gaps:
        raise HTTPException(409, "Not complete yet: " + ", ".join(gaps) + ".")

    requested = {i.document for i in application.to_gather if i.state == "uploaded"}
    documents = db.scalars(select(Document).where(Document.user_id == user.id, Document.kind.in_(requested)))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        # one plain text per item, numbered in the call's order: clean to paste into the call's form
        text_files = [f"{n:02d} {_safe_name(item.title)}.txt" for n, item in enumerate(application.to_draft, start=1)]
        for name, item in zip(text_files, application.to_draft):
            zip_file.writestr(f"texts/{name}", (item.draft or "").strip() + "\n")
        zip_file.writestr("application.docx", application_docx(call, application))
        zip_file.writestr("checklist.txt", checklist_text(call, application, text_files))
        for d in documents:
            # "cv - my_cv.pdf": the slot first, so two uploads with the same name can't collide
            zip_file.write(stored_path(d), f"{d.kind} - {os.path.basename(d.file_name)}")
        # each image item's chosen images, numbered in the artist's order, named after their project
        images = _project_images(db, user.id)
        titles = {p.id: p.title or "project" for p in _saved_projects(db, user.id)}
        for n, item in enumerate(application.to_gather, start=1):
            if item.state != "chosen":
                continue
            for m, shown in enumerate(item.images, start=1):
                image = images[str(shown.id)]
                name = f"{m:02d} {_safe_name(titles.get(image.project_id, 'project'))}.jpg"
                zip_file.write(image_path(image), f"{_image_folder(n, item)}/{name}")
    name = _safe_name(call.title_en).replace(" ", "-")
    return Response(buffer.getvalue(), media_type="application/zip",
                    headers={"Content-Disposition": f'attachment; filename="{name}.zip"'})
